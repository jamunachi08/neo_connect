"""
Frappe side of the sample analyser: reads stored samples, fills the suggested mapping, creates or
updates the API Endpoint (and optional custom fields) and validates it with a full dry run.
"""

import json
import re

import frappe
from frappe import _
from frappe import scrub
from frappe.utils import cint, now_datetime

from . import analyzer
from .utils import dumps, parse_json_field

HANDLERS = "neo_connect.handlers.commerce."
MANAGED_SETTING_KEYS = (
    "customer_fields", "customer_email_path", "customer_name_paths", "customer_phone_path", "billing_address_path", "address_map",
    "items_path", "shipping_amount_path", "discount_amount_path", "discount_on", "grand_total_path",
    "payment_method_path", "paid_amount_path", "payment_reference_path", "cost_center_path",
)


# ---------------------------------------------------------------------------
# reading samples
# ---------------------------------------------------------------------------

def read_row_text(row):
    if row.attachment:
        file_doc = frappe.get_doc("File", {"file_url": row.attachment})
        content = file_doc.get_content()
        return content.decode("utf-8-sig", errors="replace") if isinstance(content, bytes) else content
    return row.payload or ""


def load_records(sample, spec_fields=None):
    records, notes = [], []
    for row in sample.payloads:
        try:
            text = read_row_text(row)
            if spec_fields is not None:
                spec_fields += analyzer.parse_spec_tables(text)
            if "```" not in text and not text.lstrip().startswith(("{", "[")):
                row.record_count = 0
                row.parse_note = "Specification only (field tables read)" if spec_fields else "No JSON found"
                continue
            recs, row_notes = analyzer.parse_payload_text(text)
            row.record_count = len(recs)
            row.parse_note = "\n".join(row_notes) or "OK"
            if row.attachment and row.source == "Pasted":
                row.source = "Uploaded"
            records += recs
            notes += [f"Sample {row.idx}: {n}" for n in row_notes]
        except Exception as e:  # noqa: BLE001
            row.record_count = 0
            row.parse_note = f"ERROR: {e}"
            notes.append(f"Sample {row.idx}: {e}")
    return records, notes


# ---------------------------------------------------------------------------
# target description
# ---------------------------------------------------------------------------

def _fdesc(df):
    return {"fieldname": df.fieldname, "label": df.label, "fieldtype": df.fieldtype, "options": df.options,
            "reqd": cint(df.reqd), "fetch_from": df.fetch_from, "hidden": cint(df.hidden),
            "read_only": cint(df.read_only)}


def describe_target(doctype):
    meta = frappe.get_meta(doctype)
    desc = {"doctype": doctype, "fields": [_fdesc(df) for df in meta.fields if df.fieldname], "tables": {}}
    items = meta.get_field("items")
    if items and items.fieldtype in frappe.model.table_fields:
        cmeta = frappe.get_meta(items.options)
        desc["tables"]["items"] = {"doctype": items.options,
                                   "fields": [_fdesc(df) for df in cmeta.fields if df.fieldname]}
    desc["modes_of_payment"] = frappe.get_all("Mode of Payment", filters={"enabled": 1}, pluck="name")
    return desc


# ---------------------------------------------------------------------------
# analyse
# ---------------------------------------------------------------------------

def analyze(sample):
    spec_fields = []
    records, notes = load_records(sample, spec_fields)
    if not records and not spec_fields:
        frappe.throw(_("No usable JSON or spec field tables in the samples. Paste or upload at least one payload."))
    if not records:
        records = [_record_from_spec(spec_fields)]
        notes.append("No sample JSON: the structure was built from the specification tables only.")
    platform = frappe.db.get_value("API Partner", sample.partner, "platform")
    prof = analyzer.profile(records)
    spec_report = []
    if spec_fields:
        prof, spec_report = analyzer.merge_spec(prof, spec_fields)
    res = analyzer.suggest(prof, describe_target(sample.reference_doctype), platform,
                           {"payment_posting": sample.get("payment_posting")})
    allowed = {df.fieldname for df in frappe.get_meta("API Sample Field").fields}
    # rows the user reviewed/edited (or added by hand) are kept exactly as they are
    kept = [{k: r.get(k) for k in allowed} for r in sample.fields if cint(r.reviewed)]
    kept_paths = {r["external_path"] for r in kept if r.get("external_path")}
    kept_keys = {(r.get("external_path") or "", r.get("erp_field") or "") for r in kept}
    sample.set("fields", [])
    for r in kept:
        if r.get("external_path") and r["external_path"] not in prof and r.get("source") != "Manual":
            r["note"] = ("(not in the current samples) " + (r.get("note") or ""))[:500]
        sample.append("fields", r)
    for r in res["rows"]:
        r = dict(r)
        if r.get("external_path") in kept_paths or (r.get("external_path") or "", r.get("erp_field") or "") in kept_keys:
            continue
        if isinstance(r.get("value_map"), dict):
            r["value_map"] = dumps(r["value_map"])
        sample.append("fields", {k: v for k, v in r.items() if k in allowed})
    res.pop("rows")
    res["record_count"] = len(records)
    res["spec_report"] = spec_report
    sample.analysis_json = dumps(res)
    diffs = [d for d in spec_report if d["status"] != "Sample only"]
    extra = [d for d in spec_report if d["status"] == "Sample only"]
    sample.spec_differences = "\n".join(
        [f"{d['status']}: spec '{d['spec_name']}' -> JSON '{d['json_path']}'" for d in diffs]
        + [f"Sample only: '{d['json_path']}' is not documented in the spec" for d in extra]) if spec_fields else ""
    mapped = sum(1 for f in sample.fields if f.include)
    unmapped = sum(1 for f in sample.fields if f.role == "Unmapped")
    sample.analysis_notes = "\n".join(
        [f"{len(records)} sample record(s), {len(sample.fields)} fields: {mapped} used, {unmapped} unmapped."]
        + notes + res["notes"] + ([f"{len(kept)} reviewed row(s) kept as you left them."] if kept else []))
    sample.status = "Gap Ready"
    sample.payload_signature = payload_signature(sample)
    return res


def payload_signature(sample):
    """Changes whenever a sample is added, removed or edited (used to re-prepare the gap automatically)."""
    import hashlib

    h = hashlib.sha1()
    h.update((sample.reference_doctype or "").encode())
    h.update((sample.get("payment_posting") or "").encode())
    for row in sample.payloads:
        h.update((row.attachment or "").encode())
        h.update((row.payload or "").encode())
    return h.hexdigest()


def _record_from_spec(spec_fields):
    """Build a skeleton record from spec tables when no sample JSON is available."""
    from .mapper import set_path

    rec, item = {}, {}
    for f in spec_fields:
        ex = f["examples"][0] if f["examples"] else ""
        if f["type"] == "number":
            try:
                ex = float(ex)
            except ValueError:
                ex = 0.0
        if f["items"]:
            item[f["name"].split(".")[-1]] = ex
        else:
            set_path(rec, f["name"], ex)
    if item:
        rec["items"] = [item]
    return rec


# ---------------------------------------------------------------------------
# custom fields
# ---------------------------------------------------------------------------

_FIELDTYPE_BY_DATA = {"date": "Date", "datetime": "Datetime", "number": "Float", "integer": "Int",
                      "boolean": "Check"}


def create_custom_field(sample, row):
    """Create the proposed field (Fields table: proposed_* columns) on the row's Target DocType."""
    from frappe.custom.doctype.custom_field.custom_field import create_custom_field as _ccf

    path = row.external_path or ""
    in_items = "[]" in path
    prop = analyzer.propose_field(path, {"type": row.data_type or "text"}, in_items)
    fieldname = (row.proposed_fieldname or prop["proposed_fieldname"]).strip()
    if not fieldname.startswith("custom_"):
        fieldname = "custom_" + fieldname
    fieldname = scrub(fieldname)[:64]
    label = row.proposed_label or prop["proposed_label"]
    fieldtype = row.proposed_fieldtype or prop["proposed_fieldtype"]
    options = row.proposed_options if row.proposed_options is not None else prop["proposed_options"]
    meta = frappe.get_meta(sample.reference_doctype)
    items_dt = meta.get_field("items").options if meta.has_field("items") else None
    target_dt = row.target_doctype or (items_dt if in_items else sample.reference_doctype)

    if target_dt == items_dt:
        insert_after = "item_name"
    else:
        section = "custom_neo_connect_section"
        if not frappe.get_meta(target_dt).has_field(section):
            _ccf(target_dt, {"fieldname": section, "fieldtype": "Section Break",
                             "label": f"{sample.partner} Data", "collapsible": 1,
                             "insert_after": frappe.get_meta(target_dt).fields[-1].fieldname})
        insert_after = section
    if not frappe.get_meta(target_dt).has_field(fieldname):
        _ccf(target_dt, {"fieldname": fieldname, "fieldtype": fieldtype, "label": label, "options": options or None,
                         "insert_after": insert_after, "allow_on_submit": 1, "translatable": 0})

    row.proposed_fieldname = fieldname
    row.target_doctype = target_dt
    row.include = 1
    row.create_field = 0
    if target_dt == "Customer" and sample.reference_doctype != "Customer":
        row.role, row.erp_field, row.transform = "Customer Field", fieldname, "None"
        row.gap_status = "Customer Field"
    else:
        row.role = "Field"
        row.erp_field = ("items." if target_dt == items_dt else "") + fieldname
        row.transform = {"Date": "Date", "Datetime": "Datetime", "Float": "Float", "Currency": "Float",
                         "Int": "Integer", "Check": "Check"}.get(fieldtype, "Text")
        row.gap_status = "Existing Field"
    row.note = f"Created {fieldtype} field {fieldname} on {target_dt}"
    row.created_field = f"{target_dt}.{fieldname}"
    row.reviewed = 1
    created = [x for x in (sample.created_fields or "").splitlines() if x.strip()]
    if row.created_field not in created:
        created.append(row.created_field)
    sample.created_fields = "\n".join(created)
    return row.created_field


def create_ticked_fields(sample):
    return [create_custom_field(sample, row) for row in sample.fields if cint(row.create_field)]


# ---------------------------------------------------------------------------
# defaults that can be derived from the company
# ---------------------------------------------------------------------------

def derive_defaults(sample, company, plan):
    meta = frappe.get_meta(sample.reference_doctype)
    d = {}
    if meta.has_field("taxes_and_charges") and company:
        tmpl_dt = meta.get_field("taxes_and_charges").options
        tmpl = frappe.db.get_value(tmpl_dt, {"company": company, "is_default": 1, "disabled": 0}) \
            or frappe.db.get_value(tmpl_dt, {"company": company, "disabled": 0})
        if tmpl:
            d["taxes_and_charges"] = tmpl
    if meta.has_field("set_posting_time"):
        d["set_posting_time"] = 1
    if meta.has_field("disable_rounded_total"):
        d["disable_rounded_total"] = 1
    if meta.has_field("update_stock"):
        d["update_stock"] = 0
    items_dt = meta.get_field("items").options if meta.has_field("items") else None
    if items_dt and frappe.get_meta(items_dt).has_field("warehouse") and company:
        wh = frappe.db.get_single_value("Stock Settings", "default_warehouse")
        if not wh or frappe.db.get_value("Warehouse", wh, "company") != company:
            wh = frappe.db.get_value("Warehouse", {"company": company, "is_group": 0, "disabled": 0})
        if wh:
            d["items.warehouse"] = wh
    if not plan["settings"].get("cost_center_path") and company:
        cc_field = meta.get_field("cost_center")
        if cc_field and cint(cc_field.reqd):
            cc = frappe.get_cached_value("Company", company, "cost_center")
            if cc:
                d["cost_center"] = cc
    return d


# ---------------------------------------------------------------------------
# generate / update endpoint
# ---------------------------------------------------------------------------

def generate(sample):
    if not sample.analysis_json:
        analyze(sample)
    plan = json.loads(sample.analysis_json)
    partner = frappe.get_doc("API Partner", sample.partner)
    company = partner.company

    created = create_ticked_fields(sample)

    maps = []
    ext_path = None
    customer_fields = {}
    for row in sample.fields:
        if row.role == "External ID" and row.include:
            ext_path = row.external_path
        if row.role == "Customer Field":
            if row.include and row.erp_field and row.external_path:
                customer_fields[row.erp_field] = row.external_path
            continue
        if not (row.include and row.erp_field):
            continue
        if cint(row.keep_as_sent):
            fb = (row.fallback_path or "").strip()
            maps.append({
                "external_path": row.external_path, "erp_field": row.erp_field,
                "transform": "Expression" if fb else "None",
                "expression": f"first_of(record, {row.external_path!r}, {fb!r})" if fb else None,
                "add_missing_option": 1, "required": cint(row.required),
                "notes": "Kept as sent" + (f"; if empty: {fb}" if fb else ""),
            })
            continue
        maps.append({
            "external_path": row.external_path, "erp_field": row.erp_field, "transform": row.transform or "None",
            "required": cint(row.required), "lookup_doctype": row.lookup_doctype, "lookup_field": row.lookup_field,
            "if_not_found": row.if_not_found or "Error", "source_timezone": row.source_timezone,
            "value_map": row.value_map, "expression": row.expression, "notes": (row.note or "")[:140],
        })
    ext_path = ext_path or plan.get("external_id_path")
    if not ext_path:
        frappe.throw(_("No External ID: set the Role 'External ID' on the field that holds the unique order number."))

    settings = dict(plan["settings"])
    settings.pop("customer_fields", None)
    if customer_fields:
        settings["customer_fields"] = customer_fields
        settings.setdefault("update_customer_fields", 1)
    pre = [HANDLERS + h for h in plan["pre_hooks"]]
    if customer_fields and HANDLERS + "ensure_customer" not in pre:
        pre.insert(0, HANDLERS + "ensure_customer")
    post = [HANDLERS + h for h in plan["post_hooks"]]
    defaults = derive_defaults(sample, company, plan)

    slug = (sample.endpoint_slug or f"{scrub(sample.partner)}-{scrub(sample.reference_doctype)}")
    slug = re.sub(r"[^a-z0-9]+", "-", slug.lower()).strip("-")
    records_total = plan.get("record_count") or 0
    first_record = _first_record(sample)

    if sample.endpoint and frappe.db.exists("API Endpoint", sample.endpoint):
        ep = frappe.get_doc("API Endpoint", sample.endpoint)
        old_defaults = parse_json_field(ep.defaults, {}) or {}
        defaults.update(old_defaults)                      # the user's own defaults win
        old_settings = parse_json_field(ep.handler_settings, {}) or {}
        for k, v in old_settings.items():
            if k in MANAGED_SETTING_KEYS:
                continue
            if k not in settings and "<" in json.dumps(v, ensure_ascii=False):
                continue                                   # unfilled placeholder that is no longer needed
            settings[k] = v                                # keep the user's business settings
        if plan.get("confirmed_mop") and plan["settings"].get("mode_of_payment_map"):
            settings["mode_of_payment_map"] = plan["settings"]["mode_of_payment_map"]   # confirmed in Review
        pre = _merge_lines(pre, ep.pre_process_hook)
        post = _merge_lines(post, ep.post_process_hook)
        ep.set("field_maps", [])
    else:
        if frappe.db.exists("API Endpoint", slug):
            frappe.throw(_("Endpoint {0} already exists. Set a different Endpoint Code.").format(slug))
        ep = frappe.new_doc("API Endpoint")
        ep.slug = slug
        ep.endpoint_title = sample.title
        ep.direction = "Inbound"
        ep.enabled = 1
        ep.on_duplicate = "Skip"
        ep.append("allowed_partners", {"partner": sample.partner})

    ep.reference_doctype = sample.reference_doctype
    ep.external_id_path = ext_path
    ep.document_action = sample.document_action or "Save as Draft"
    ep.description = ep.description or f"Generated from Payload Sample {sample.name} ({records_total} sample record(s))."
    ep.pre_process_hook = "\n".join(pre)
    ep.post_process_hook = "\n".join(post)
    ep.defaults = dumps(defaults)
    ep.handler_settings = dumps(settings)
    ep.sample_payload = dumps(first_record) if first_record else ep.sample_payload
    for m in maps:
        ep.append("field_maps", m)
    ep.save() if not ep.is_new() else ep.insert()

    sample.endpoint = ep.name
    sample.status = "Endpoint Generated"
    sample.validation_report = validate(sample, ep, created, settings)
    return ep.name


def _merge_lines(generated, existing):
    out = list(generated)
    for line in (existing or "").splitlines():
        if line.strip() and line.strip() not in out:
            out.append(line.strip())
    return out


def _first_record(sample):
    for row in sample.payloads:
        try:
            recs, _n = analyzer.parse_payload_text(read_row_text(row))
            if recs:
                return recs[0]
        except Exception:  # noqa: BLE001
            continue
    return None


# ---------------------------------------------------------------------------
# validation report
# ---------------------------------------------------------------------------

_MANDATORY_RE = re.compile(r"\[([^,\]]+),\s*[^\]]*\]:\s*(.+)")


def validate(sample, ep, created, settings):
    from .inbound import Processor

    lines = [f"Endpoint: {ep.name}", f"URL: {ep.endpoint_url}", ""]
    if created:
        lines.append("Custom fields created: " + ", ".join(created))
    todo = []
    for key, val in settings.items():
        if "<" in json.dumps(val, ensure_ascii=False):
            todo.append(f"Handler Settings '{key}': replace the placeholder values (in angle brackets)")
    for row in sample.fields:
        if row.include and row.value_map and "<" in row.value_map and not cint(row.keep_as_sent):
            todo.append(f"Field {row.erp_field}: fill the Value Map with the allowed options")
    defaults = parse_json_field(ep.defaults, {}) or {}
    if frappe.get_meta(ep.reference_doctype).has_field("taxes_and_charges") and not defaults.get("taxes_and_charges"):
        todo.append("Document Defaults: no Sales Taxes template found for the company; set 'taxes_and_charges'")

    partner = frappe.get_doc("API Partner", sample.partner)
    results = []
    records = []
    for row in sample.payloads:
        try:
            records += analyzer.parse_payload_text(read_row_text(row))[0]
        except Exception:  # noqa: BLE001
            pass
    for i, rec in enumerate(records[:5]):
        sp = f"cc_validate_{i}"
        frappe.db.savepoint(sp)
        try:
            out = Processor(ep, partner).process_payload(rec, dry_run=True)
            results.append(out["results"][0])
        except Exception as e:  # noqa: BLE001
            results.append({"status": "Failed", "error": str(e)})
        finally:
            frappe.db.rollback(save_point=sp)
            frappe.clear_messages()

    lines.append("Dry run with the sample payload(s):")
    missing = {}
    for i, r in enumerate(results, 1):
        lines.append(f"  Sample {i}: {r.get('status')}" + (f" - {r.get('error')}" if r.get("error") else ""))
        if " not found for " in (r.get("error") or ""):
            lines.append("     fix: create that record in ERPNext, translate it with a Value Map, "
                         "or set the row's 'If Not Found' to 'Use Value'/'Skip'")
        for w in r.get("warnings") or []:
            lines.append(f"     warning: {w}")
        m = _MANDATORY_RE.search(r.get("error") or "")
        if m:
            for fn in [x.strip() for x in m.group(2).split(",") if x.strip()]:
                missing.setdefault(fn, m.group(1).strip())

    if missing:
        lines += ["", "Mandatory fields still without a value (from ERPNext's own validation):"]
        meta = frappe.get_meta(ep.reference_doctype)
        items_dt = meta.get_field("items").options if meta.has_field("items") else None
        for fn in missing:
            df = meta.get_field(fn)
            where = ""
            if not df and items_dt:
                df = frappe.get_meta(items_dt).get_field(fn)
                where = "items."
            label = (df.label if df else fn) or fn
            hint = ""
            if df and df.fieldtype == "Select":
                hint = "options: " + " | ".join(o for o in (df.options or "").split("\n") if o)
            elif df and df.fieldtype == "Link":
                hint = f"a {df.options} name"
            lines.append(f"  - {where}{fn} ({label}): {hint}")
            mapped = next((r for r in sample.fields if r.include and r.erp_field == where + fn), None)
            if mapped and mapped.transform == "Value Map":
                lines.append(f"      fix: fill the Value Map of {where}{fn} (Fields table) so every payload value "
                             f"becomes one of the options, then Update Endpoint")
            elif mapped:
                lines.append(f"      mapped from '{mapped.external_path}' (present in {mapped.presence} samples): "
                             f"add \"{where}{fn}\": \"VALUE\" to Document Defaults for orders without it")
            else:
                cand = _best_candidate(sample, df) if df else None
                lines.append(f"      fix: add \"{where}{fn}\": \"VALUE\" to Document Defaults"
                             + (f", or map it from '{cand}' in the Fields table" if cand else ""))
    if todo:
        lines += ["", "To do before go-live:"] + [f"  - {t}" for t in todo]
    if results and all(r.get("status") == "Validated" for r in results) and not todo:
        lines += ["", "READY: all samples validate. Share the Partner Guide from the endpoint."]
    return "\n".join(lines)


def _best_candidate(sample, df):
    best, score = None, 0
    for row in sample.fields:
        if row.external_path and row.role in ("Unmapped", "Ignore"):
            s = analyzer.match_score(row.external_path, {"fieldname": df.fieldname, "label": df.label})
            if s > score:
                best, score = row.external_path, s
    return best if score >= 0.3 else None


# ---------------------------------------------------------------------------
# samples pushed by partners through the API
# ---------------------------------------------------------------------------

def store_partner_sample(partner, message_type, raw_text, label=None):
    analyzer.parse_payload_text(raw_text)   # raises if not JSON
    name = frappe.db.get_value("API Payload Sample", {"partner": partner, "message_type": message_type})
    if name:
        sample = frappe.get_doc("API Payload Sample", name)
    else:
        ref_dt = message_type if frappe.db.exists("DocType", message_type) else "Sales Order"
        sample = frappe.get_doc({"doctype": "API Payload Sample", "title": f"{partner} - {message_type}",
                                 "partner": partner, "message_type": message_type, "reference_doctype": ref_dt})
    sample.append("payloads", {"label": label or f"API {now_datetime():%Y-%m-%d %H:%M}", "payload": raw_text,
                               "source": "Partner API", "received_on": now_datetime()})
    sample.flags.ignore_permissions = True
    sample.flags.from_partner_api = True          # gap is prepared automatically on save
    sample.save() if not sample.is_new() else sample.insert()
    return sample


# ---------------------------------------------------------------------------
# gap report (screen + Excel)
# ---------------------------------------------------------------------------

_AUTO_FILLED = {
    "naming_series", "company", "currency", "conversion_rate", "selling_price_list", "buying_price_list",
    "price_list_currency", "plc_conversion_rate", "debit_to", "credit_to", "customer", "supplier", "items",
    "item_name", "uom", "stock_uom", "conversion_factor", "amount", "base_rate", "base_amount", "income_account",
    "expense_account", "description", "posting_date", "transaction_date", "base_net_total", "base_grand_total",
    "grand_total", "net_total", "status", "title",
}


def gap_report(sample):
    if not sample.fields:
        frappe.throw(_("Click Analyze first"))
    plan = json.loads(sample.analysis_json or "{}")
    ep_defaults = {}
    if sample.endpoint and frappe.db.exists("API Endpoint", sample.endpoint):
        ep_defaults = parse_json_field(frappe.db.get_value("API Endpoint", sample.endpoint, "defaults"), {}) or {}
    F = sample.fields

    def rowval(r, *keys):
        return [r.get(k) or "" for k in keys]

    sections = []
    create = [r for r in F if r.gap_status == "Create Field"]
    sections.append({
        "title": "Fields to create", "key": "create",
        "hint": "Values the partner sends that have no field in ERPNext. Tick 'Create Field' in the Fields table "
                "for the ones to keep, then click 'Create Ticked Fields' or Generate.",
        "columns": ["Target DocType", "Fieldname", "Label", "Field Type", "Options", "Partner JSON Path", "Sample",
                    "In Samples", "Source", "Ticked", "Note"],
        "rows": [rowval(r, "target_doctype", "proposed_fieldname", "proposed_label", "proposed_fieldtype",
                        "proposed_options", "external_path", "sample_value", "presence", "source")
                 + ["Yes" if r.create_field else "", r.note or ""] for r in create],
    })
    opt = []
    for r in F:
        if r.gap_status != "Option Gap":
            continue
        vmap = parse_json_field(r.value_map, {}) or {}
        missing = [k for k, v in vmap.items() if k != "*" and str(v).startswith("<")]
        field = frappe.get_meta(r.target_doctype).get_field(r.erp_field.split(".")[-1]) if r.target_doctype else None
        options = " | ".join(o for o in ((field.options if field else "") or "").split("\n") if o)
        opt.append([r.target_doctype, r.erp_field, r.external_path, ", ".join(missing), options,
                    "Add the missing options in Customize Form, or translate them in the Value Map"])
    sections.append({"title": "Select option gaps", "key": "options",
                     "hint": "Payload values with no matching option in a Select field.",
                     "columns": ["Target DocType", "Field", "Partner JSON Path", "Values without option",
                                 "Existing options", "Action"], "rows": opt})
    sections.append({"title": "Written to the Customer", "key": "customer",
                     "hint": "The invoice fetches these from the Customer, so NeoConnect updates the Customer.",
                     "columns": ["Customer Field", "Partner JSON Path", "Sample", "Note"],
                     "rows": [rowval(r, "erp_field", "external_path", "sample_value", "note") for r in F
                              if r.role == "Customer Field"]})
    sections.append({"title": "Mapped to existing fields", "key": "existing",
                     "columns": ["Target DocType", "ERPNext Field", "Partner JSON Path", "Transform", "Sample", "Use"],
                     "rows": [rowval(r, "target_doctype", "erp_field", "external_path", "transform", "sample_value")
                              + ["Yes" if r.include else "No"] for r in F if r.gap_status == "Existing Field"]})

    # mandatory fields without any source
    mapped = {r.erp_field for r in F if r.include and r.erp_field}
    meta = frappe.get_meta(sample.reference_doctype)
    items_dt = meta.get_field("items").options if meta.has_field("items") else None
    mand = []
    for dt_, prefix in ((sample.reference_doctype, ""), (items_dt, "items.")):
        if not dt_:
            continue
        for df in frappe.get_meta(dt_).fields:
            if not cint(df.reqd) or df.fieldtype in frappe.model.table_fields or df.fieldname in _AUTO_FILLED \
                    or cint(df.read_only) or df.default not in (None, ""):
                continue
            key = prefix + df.fieldname
            if key in mapped or key in ep_defaults or (df.fieldname == "cost_center" and
                                                       plan.get("settings", {}).get("cost_center_path")):
                continue
            hint = ""
            if df.fieldtype == "Select":
                hint = "options: " + " | ".join(o for o in (df.options or "").split("\n") if o)
            elif df.fieldtype == "Link":
                hint = f"a {df.options}"
            mand.append([dt_, key, df.label or "", df.fieldtype, hint,
                         "Map it in the Fields table, or set a Document Default on the endpoint"])
    sections.append({"title": "Mandatory fields without a source", "key": "mandatory",
                     "hint": "Required by your ERPNext but not filled by the mapping or a default yet.",
                     "columns": ["DocType", "Field", "Label", "Type", "Expected", "Action"], "rows": mand})
    sections.append({"title": "Specification vs sample JSON", "key": "spec",
                     "hint": "Send these to the partner: names that differ, or fields only one of them has.",
                     "columns": ["Status", "Name in spec", "Path in JSON", "Note"],
                     "rows": [[d["status"], d["spec_name"], d["json_path"], d["note"]]
                              for d in plan.get("spec_report") or []]})
    sections.append({"title": "Handled by NeoConnect / not needed", "key": "handled",
                     "columns": ["Partner JSON Path", "Role", "Status", "Note"],
                     "rows": [rowval(r, "external_path", "role", "gap_status", "note") for r in F
                              if r.gap_status in ("Handled by NeoConnect", "Calculated / Not needed")]})
    summary = {s_["title"]: len(s_["rows"]) for s_ in sections}
    return {"sample": sample.name, "title": sample.title, "target": sample.reference_doctype,
            "summary": summary, "sections": sections}


def gap_report_html(rep):
    from frappe.utils import escape_html as esc

    parts = [f"<p><b>{esc(rep['title'])}</b> &rarr; {esc(rep['target'])}</p><p>"]
    parts.append(" &nbsp; ".join(f"<span class='indicator-pill {'red' if n and k in ('Fields to create', 'Select option gaps', 'Mandatory fields without a source') else 'green' if not n else 'blue'}'>"
                                 f"{esc(k)}: {n}</span>" for k, n in rep["summary"].items()))
    parts.append("</p>")
    for sec in rep["sections"]:
        parts.append(f"<h5 style='margin-top:18px'>{esc(sec['title'])} ({len(sec['rows'])})</h5>")
        if sec.get("hint"):
            parts.append(f"<p class='text-muted small'>{esc(sec['hint'])}</p>")
        if not sec["rows"]:
            parts.append("<p class='text-muted small'>None</p>")
            continue
        parts.append("<div style='overflow-x:auto'><table class='table table-bordered table-sm small'><thead><tr>")
        parts += [f"<th>{esc(c)}</th>" for c in sec["columns"]]
        parts.append("</tr></thead><tbody>")
        for r in sec["rows"]:
            parts.append("<tr>" + "".join(f"<td>{esc(str(v))}</td>" for v in r) + "</tr>")
        parts.append("</tbody></table></div>")
    return "".join(parts)


def gap_report_xlsx(rep):
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append([f"Gap report: {rep['title']} -> {rep['target']}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    for k, n in rep["summary"].items():
        ws.append([k, n])
    ws.column_dimensions["A"].width = 42
    head_fill = PatternFill("solid", fgColor="DDE7F3")
    for sec in rep["sections"]:
        sh = wb.create_sheet(re.sub(r"[\\/*?:\[\]]", "", sec["title"])[:31])
        sh.append(sec["columns"])
        for c in sh[1]:
            c.font = Font(bold=True)
            c.fill = head_fill
        for r in sec["rows"]:
            sh.append([str(v) if v is not None else "" for v in r])
        for i, col in enumerate(sec["columns"], 1):
            width = max([len(str(col))] + [len(str(r[i - 1])) for r in sec["rows"] if i - 1 < len(r)] + [8])
            sh.column_dimensions[sh.cell(1, i).column_letter].width = min(width + 2, 60)
        for row in sh.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        sh.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# review & confirm
# ---------------------------------------------------------------------------

def _options_of(doctype, fieldname):
    if not doctype or not fieldname:
        return []
    df = frappe.get_meta(doctype).get_field(fieldname.split(".")[-1])
    if not df or df.fieldtype != "Select":
        return []
    return [o for o in (df.options or "").split("\n") if o.strip()]


def review_data(sample):
    """Everything the Review & Confirm screen needs, grouped."""
    meta = frappe.get_meta(sample.reference_doctype)
    items_dt = meta.get_field("items").options if meta.has_field("items") else None
    targets = [sample.reference_doctype] + ([items_dt] if items_dt else []) + ["Customer"]
    create, mapped, options, other = [], [], [], []
    for r in sample.fields:
        base = {"name": r.name, "path": r.external_path or "", "sample": r.sample_value or "", "presence": r.presence or "",
                "source": r.source or "", "include": cint(r.include), "note": r.note or "", "reviewed": cint(r.reviewed)}
        if r.gap_status == "Create Field" or (r.role == "Unmapped" and not r.created_field):
            prop = analyzer.propose_field(r.external_path or "x", {"type": r.data_type or "text"}, "[]" in (r.external_path or ""))
            create.append(dict(base, create=cint(r.create_field), target=r.target_doctype or sample.reference_doctype,
                               fieldname=r.proposed_fieldname or prop["proposed_fieldname"],
                               label=r.proposed_label or prop["proposed_label"],
                               fieldtype=r.proposed_fieldtype or prop["proposed_fieldtype"],
                               options=r.proposed_options or ""))
        elif r.gap_status == "Option Gap" or (r.transform == "Value Map" and r.erp_field):
            opts = _options_of(r.target_doctype, r.erp_field)
            vmap = parse_json_field(r.value_map, {}) or {}
            options.append(dict(base, field=r.erp_field, target=r.target_doctype, select_options=opts,
                                keep_as_sent=cint(r.keep_as_sent), fallback=r.fallback_path or "",
                                values=[{"value": k, "mapped": v if v in opts else ""} for k, v in vmap.items()]))
        elif r.gap_status in ("Existing Field", "Customer Field") and r.erp_field:
            mapped.append(dict(base, field=r.erp_field, target=r.target_doctype or "", transform=r.transform or "None",
                               role=r.role))
        else:
            other.append(dict(base, role=r.role, status=r.gap_status))
    fieldtypes = [o for o in (frappe.get_meta("API Sample Field").get_field("proposed_fieldtype").options or "").split("\n") if o]
    plan = json.loads(sample.analysis_json or "{}")
    mop_map = (plan.get("settings") or {}).get("mode_of_payment_map") or {}
    mops = frappe.get_all("Mode of Payment", filters={"enabled": 1}, pluck="name", order_by="name")
    payments = [{"value": k, "mapped": v if v in mops else ""} for k, v in mop_map.items()]
    company = frappe.db.get_value("API Partner", sample.partner, "company")
    accounts = frappe.get_all("Account", filters={"company": company, "is_group": 0,
                                                  "account_type": ["in", ["Bank", "Cash", "Receivable"]]},
                              pluck="name", order_by="account_type, name") if company else []
    text_paths = sorted(r.external_path for r in sample.fields
                        if r.external_path and "[]" not in r.external_path and (r.data_type or "text") == "text")
    return {"targets": targets, "fieldtypes": fieldtypes, "create": create, "mapped": mapped, "options": options,
            "other": other, "status": sample.status, "endpoint": sample.endpoint, "text_paths": text_paths,
            "payments": payments, "modes_of_payment": mops, "accounts": accounts, "company": company,
            "payment_path": (plan.get("settings") or {}).get("payment_method_path")}


def apply_review(sample, decisions):
    """
    decisions = {
      "create":  [{"name", "create", "target", "fieldname", "label", "fieldtype", "options"}],
      "mapped":  [{"name", "include", "field"}],
      "options": [{"name", "map": {payload value: option}}],
      "add":     [{"path", "target", "fieldname", "label", "fieldtype", "options"}],   # user-added fields
      "generate": 1
    }
    """
    rows = {r.name: r for r in sample.fields}
    for d in decisions.get("create") or []:
        r = rows.get(d.get("name"))
        if not r:
            continue
        r.create_field = cint(d.get("create"))
        r.include = r.create_field
        r.target_doctype = d.get("target") or r.target_doctype
        r.proposed_fieldname = (d.get("fieldname") or r.proposed_fieldname or "").strip()
        r.proposed_label = d.get("label") or r.proposed_label
        r.proposed_fieldtype = d.get("fieldtype") or r.proposed_fieldtype
        r.proposed_options = d.get("options") if d.get("options") is not None else r.proposed_options
        r.reviewed = 1
    for d in decisions.get("mapped") or []:
        r = rows.get(d.get("name"))
        if not r:
            continue
        r.include = cint(d.get("include"))
        if d.get("field"):
            r.erp_field = d["field"].strip()
        r.reviewed = 1
    from .options import ensure_select_option

    for d in decisions.get("options") or []:
        r = rows.get(d.get("name"))
        if not r:
            continue
        field = (r.erp_field or "").split(".")[-1]
        typed = list(d.get("new_options") or [])
        if not cint(d.get("keep_as_sent")):
            typed += [v for k, v in (d.get("map") or {}).items() if v and not str(v).startswith("<")]
        for opt in typed:                                   # options typed by the user (e.g. تابي) are added
            if r.target_doctype:
                ensure_select_option(r.target_doctype, field, str(opt).strip())
        if cint(d.get("keep_as_sent")):
            r.keep_as_sent = 1
            r.fallback_path = (d.get("fallback") or "").strip()
            r.include = 1
            r.gap_status = "Existing Field"
            r.note = "Kept as sent; new values are added to the options automatically"
            r.reviewed = 1
            continue
        r.keep_as_sent = 0
        vmap = {k: v for k, v in (d.get("map") or {}).items() if v not in (None, "")}
        old = parse_json_field(r.value_map, {}) or {}
        for k in old:
            vmap.setdefault(k, old[k])
        r.value_map = dumps(vmap)
        r.transform = "Value Map"
        r.include = 1
        missing = [k for k, v in vmap.items() if str(v).startswith("<")]
        r.gap_status = "Option Gap" if missing else "Existing Field"
        r.reviewed = 1
    for d in decisions.get("add") or []:
        path = (d.get("path") or "").strip()
        if not path:
            continue
        sample.append("fields", {
            "include": 1, "create_field": 1, "role": "Unmapped", "gap_status": "Create Field", "source": "Manual",
            "external_path": path, "data_type": {"Date": "date", "Datetime": "datetime", "Int": "integer",
                                                 "Float": "number", "Currency": "number", "Check": "boolean"}
            .get(d.get("fieldtype"), "text"),
            "target_doctype": d.get("target") or sample.reference_doctype, "proposed_fieldname": d.get("fieldname"),
            "proposed_label": d.get("label"), "proposed_fieldtype": d.get("fieldtype") or "Data",
            "proposed_options": d.get("options"), "reviewed": 1, "note": "Added by user"})

    company = frappe.db.get_value("API Partner", sample.partner, "company")
    for nm in decisions.get("new_mop") or []:
        mop_name = (nm.get("name") or "").strip()
        if not mop_name:
            continue
        if not frappe.db.exists("Mode of Payment", mop_name):
            mop = frappe.get_doc({"doctype": "Mode of Payment", "mode_of_payment": mop_name,
                                  "type": nm.get("type") or "Bank", "enabled": 1})
            if company and nm.get("account"):
                mop.append("accounts", {"company": company, "default_account": nm["account"]})
            mop.insert(ignore_permissions=True)
        decisions.setdefault("mop", {})[nm.get("value")] = mop_name

    if decisions.get("mop"):
        plan = json.loads(sample.analysis_json or "{}")
        plan.setdefault("settings", {})["mode_of_payment_map"] = {
            k: (v or "<Mode of Payment>") for k, v in decisions["mop"].items()}
        plan["confirmed_mop"] = 1
        sample.analysis_json = dumps(plan)

    created = create_ticked_fields(sample)
    sample.confirmed_by = frappe.session.user
    sample.confirmed_on = frappe.utils.now_datetime()
    sample.status = "Confirmed"
    report = ""
    if cint(decisions.get("generate", 1)):
        generate(sample)
        report = sample.validation_report
    return {"created": created, "report": report, "endpoint": sample.endpoint}


def remove_created_fields(sample):
    """Undo: delete the custom fields this sample created (data already stored in them is not migrated)."""
    removed = []
    for ref in [x.strip() for x in (sample.created_fields or "").splitlines() if x.strip()]:
        dt, fieldname = ref.split(".", 1)
        name = frappe.db.get_value("Custom Field", {"dt": dt, "fieldname": fieldname})
        if name:
            frappe.delete_doc("Custom Field", name, ignore_permissions=True)
            removed.append(ref)
    for r in sample.fields:
        if r.created_field in removed:
            r.created_field = None
            r.role, r.gap_status, r.include, r.erp_field = "Unmapped", "Create Field", 0, None
            r.note = "Field removed (undo); create it again from Review & Confirm if needed"
    sample.created_fields = ""
    for dt in {x.split(".", 1)[0] for x in removed}:
        frappe.clear_cache(doctype=dt)
    return removed
